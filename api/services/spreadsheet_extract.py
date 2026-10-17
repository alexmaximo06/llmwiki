"""Spreadsheet parsing into per-sheet markdown tables."""

import csv

import openpyxl

MAX_TABLE_ROWS = 100


def parse_sheets(path: str, ext: str) -> list[tuple[str, str]]:
    """Return (sheet_name, markdown_table) per sheet."""
    if ext == "csv":
        with open(path, newline="", encoding="utf-8", errors="replace") as f:
            rows = [list(row) for row in csv.reader(f)]
        return [("Sheet1", _rows_to_markdown(rows))]

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = [
            [str(cell) if cell is not None else "" for cell in row]
            for row in ws.iter_rows(values_only=True)
        ]
        if not rows:
            continue
        sheets.append((name, _rows_to_markdown(rows)))
    wb.close()
    return sheets


def _rows_to_markdown(rows: list[list[str]], max_rows: int = MAX_TABLE_ROWS) -> str:
    if not rows:
        return "(empty)"
    header = "| " + " | ".join(rows[0]) + " |"
    sep = "| " + " | ".join("---" for _ in rows[0]) + " |"
    data = rows[1:max_rows + 1]
    body = "\n".join("| " + " | ".join(r) + " |" for r in data)
    truncated = f"\n\n*({len(rows) - 1 - max_rows} more rows truncated)*" if len(rows) - 1 > max_rows else ""
    return f"{header}\n{sep}\n{body}{truncated}"
